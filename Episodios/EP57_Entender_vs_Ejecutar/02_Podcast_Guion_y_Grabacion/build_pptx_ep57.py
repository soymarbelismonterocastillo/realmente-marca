import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # BRAND PALETTE: Clean high-contrast dark mode
    c_bg = RGBColor(25, 6, 51)           # #190633 Dark Purple
    c_card = RGBColor(44, 12, 85)        # #2C0C55 Card Container
    c_teal = RGBColor(42, 210, 197)      # #2AD2C5 Bright Cyan Accent
    c_pink = RGBColor(255, 90, 141)      # #FF5A8D Bright Pink Accent
    c_white = RGBColor(255, 255, 255)    # Pure White for all body copy

    def apply_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_bg
        bg.line.fill.background()

    def add_header(slide, tag, title, subtitle=None):
        tb_tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        p_tag = tb_tag.text_frame.paragraphs[0]
        p_tag.text = tag.upper()
        p_tag.font.size = Pt(14)
        p_tag.font.bold = True
        p_tag.font.color.rgb = c_teal

        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(32)
        p_title.font.bold = True
        p_title.font.color.rgb = c_white

        if subtitle:
            tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.5))
            p_sub = tb_sub.text_frame.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.size = Pt(20)
            p_sub.font.color.rgb = c_white

    # SLIDE 1: PORTADA
    s1 = prs.slides.add_slide(blank_layout)
    apply_bg(s1)
    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(4.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.vertical_anchor = MSO_ANCHOR.MIDDLE

    p1 = tf1.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    p1.text = "REALMENTE CLUB · SESIÓN EN VIVO"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = c_teal
    p1.space_after = Pt(16)

    p2 = tf1.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Tu vida no cambia cuando lo entiendes,\ncambia cuando lo ejecutas"
    p2.font.size = Pt(38)
    p2.font.bold = True
    p2.font.color.rgb = c_white
    p2.space_after = Pt(16)

    p3 = tf1.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.text = "De la teoría a la acción: Desmantelando los 4 patrones de saboteo"
    p3.font.size = Pt(22)
    p3.font.bold = True
    p3.font.color.rgb = c_teal

    tb1_foot = s1.shapes.add_textbox(Inches(1.0), Inches(6.4), Inches(11.333), Inches(0.6))
    p_foot = tb1_foot.text_frame.paragraphs[0]
    p_foot.alignment = PP_ALIGN.CENTER
    p_foot.text = "Facilitadora: Marbelis Montero  |  Sábado 19 de Septiembre, 2026"
    p_foot.font.size = Pt(16)
    p_foot.font.color.rgb = c_white

    # SLIDE 2: FRASE APERTURA
    s2 = prs.slides.add_slide(blank_layout)
    apply_bg(s2)
    add_header(s2, "REFLEXIÓN DE APERTURA", "Tu Verdadero Punto de Partida Hoy")

    qbox2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(2.3), Inches(10.9), Inches(4.0))
    qbox2.fill.solid()
    qbox2.fill.fore_color.rgb = c_card
    qbox2.line.color.rgb = c_teal
    qbox2.line.width = Pt(2)
    tf2 = qbox2.text_frame
    tf2.word_wrap = True
    tf2.vertical_anchor = MSO_ANCHOR.MIDDLE

    p = tf2.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "“Saber lo que tienes que hacer y no hacerlo es exactamente lo mismo que no saber nada. Tu vida no cambia con la teoría que entiendes en la cabeza, cambia con la imperfección que te atreves a ejecutar en el mundo real.”"
    p.font.size = Pt(28)
    p.font.italic = True
    p.font.color.rgb = c_white

    # SLIDE 3: ENCUADRE
    s3 = prs.slides.add_slide(blank_layout)
    apply_bg(s3)
    add_header(s3, "CONCEPTO CLAVE", "La Brecha Entre Entender y Ejecutar", "Por qué saberlo todo no está cambiando tu vida")

    cards3 = [
        ("📚 Zona de Confort del Saber", "Creer que acumular información, libros, podcasts y diagnósticos equivale a transformar tu vida real.", c_teal),
        ("🧠 La Ilusión del Ego", "El ego te hace sentir que 'estás trabajando en ti' al analizar, para evitar el riesgo de actuar.", c_pink),
        ("⚡ La Verdad Incómoda", "Entender calma la culpa, pero solo la ejecución concreta reprograma tu mente real.", c_teal)
    ]
    for i, (title, desc, color) in enumerate(cards3):
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.9), Inches(2.4), Inches(3.6), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = color
        card.line.width = Pt(2)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = c_white
        p.space_after = Pt(14)
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(19)
        p2.font.color.rgb = c_white

    # SLIDE 4: MAPA DEL SABOTEO
    s4 = prs.slides.add_slide(blank_layout)
    apply_bg(s4)
    add_header(s4, "MAPA MENTAL DE LA PARÁLISIS", "Las 2 Grandes Trampas Invisibles del Ego", "De dónde nace tu procrastinación sofisticada")

    cards4 = [
        ("🔴 La Trampa del Pasado", "La Adicción al 'Por Qué'", "Creer que necesitas resolver o entender hasta el último origen histórico de tu pasado antes de dar un solo paso en el presente.", c_pink),
        ("🟣 La Trampa del Futuro", "La Ilusión del 'Momento Perfecto'", "Creer que debes esperar a que tus emociones, tu seguridad, tu motivación y las circunstancias estén totalmente impecables.", c_teal)
    ]
    for i, (tag, title, desc, color) in enumerate(cards4):
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 5.9), Inches(2.4), Inches(5.6), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = color
        card.line.width = Pt(2)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p1 = tf.paragraphs[0]
        p1.text = tag
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = color
        p1.space_after = Pt(8)
        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = c_white
        p2.space_after = Pt(14)
        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(19)
        p3.font.color.rgb = c_white

    # HELPER FOR PATTERN SLIDES 5 to 8
    def build_pattern_slide(slide_num, tag, title, points):
        s = prs.slides.add_slide(blank_layout)
        apply_bg(s)
        add_header(s, tag, title)
        for i, (ptitle, pdesc, color) in enumerate(points):
            box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.1 + i * 1.6), Inches(11.7), Inches(1.35))
            box.fill.solid()
            box.fill.fore_color.rgb = c_card
            box.line.color.rgb = color
            box.line.width = Pt(2)
            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.text = ptitle
            p.font.size = Pt(22)
            p.font.bold = True
            p.font.color.rgb = color
            p.space_after = Pt(4)
            p2 = tf.add_paragraph()
            p2.text = pdesc
            p2.font.size = Pt(19)
            p2.font.color.rgb = c_white

    # SLIDE 5
    build_pattern_slide(5, "TRAMPA DEL PASADO · PATRÓN 1", "La Analista del Pasado ('La Adicción al Por Qué')", [
        ("🔍 El Mecanismo del Ego", "Creer que comprender la causa original o el trauma del pasado equivale automáticamente a haber resuelto el problema.", c_pink),
        ("🔄 El Bucle de las Explicaciones", "Buscar infinitos justificativos en tu historia para darte la razón y quedarte en la zona conocida.", c_teal),
        ("💥 La Verdad Incómoda", "Entender a medias consuela la culpa del ego, pero no cambia tu comportamiento presente si no ejecutas.", c_white)
    ])

    # SLIDE 6
    build_pattern_slide(6, "TRAMPA DEL PASADO · PATRÓN 2", "La Sanadora Eterna ('La Excusa del Trabajo Previo')", [
        ("🛡️ El Escudo de la Sanación", "Usar la sanación como una coraza sofisticada para evitar la exposición, el rechazo y el riesgo de actuar.", c_pink),
        ("🚫 La Premisa Limítrofe", "Creer firmemente: 'Hasta que no sane esto del todo o termine este taller, no me puedo exigir avanzar'.", c_white),
        ("✨ La Verdad Incómoda", "La acción concreta en el presente es la mejor herramienta de reprogramación y sanación, no el requisito previo.", c_teal)
    ])

    # SLIDE 7
    build_pattern_slide(7, "TRAMPA DEL FUTURO · PATRÓN 3", "La Esperadora del Momento Perfecto ('La Ilusión del Futuro')", [
        ("⏳ El Mecanismo de Postergación", "Postergar decisiones esperando que las circunstancias o tu estado emocional estén mágicamente alineados.", c_teal),
        ("🔮 El Mito de Sentirse Lista", "Fantasear con que un día te despertarás sintiéndote 100% segura, confiada, sin dudas y sin ningún tipo de miedo.", c_pink),
        ("🎯 La Verdad Incómoda", "El momento perfecto no existe. Esperarlo es una defensa contra la incomodidad de aprender sobre la marcha.", c_white)
    ])

    # SLIDE 8
    build_pattern_slide(8, "TRAMPA DEL FUTURO · PATRÓN 4", "La Protectora del Potencial ('La Fantasía del Control')", [
        ("💎 La Fantasía Idealizada", "Enamorarse de las ideas en la cabeza, donde todo es impecable porque nadie las juzga ni las pone a prueba.", c_teal),
        ("🫣 El Miedo al Choque de Realidad", "Aterra chocar con la realidad concreta porque la realidad siempre es imperfecta, caótica e incontrolable.", c_pink),
        ("🔨 La Verdad Incómoda", "Es mil veces preferible un primer intento imperfecto que conservar el estatus de 'alguien con gran potencial no realizado'.", c_white)
    ])

    # SLIDE 9: NADAR
    s9 = prs.slides.add_slide(blank_layout)
    apply_bg(s9)
    add_header(s9, "APRENDIZAJE EN ACCIÓN", "La Metáfora de Lanzarse al Agua", "Aprender a nadar exige estar en el medio")
    cards9 = [
        ("🏖️ Fuera del Agua (Teoría)", "En la orilla todo es planificación limpia y perfecta. Pero afuera NO estás aprendiendo a nadar.", c_teal),
        ("🌊 Dentro del Agua (Acción)", "Para aprender a nadar hay que lanzarse, sentir la densidad del agua, patalear y tragar agua.", c_pink),
        ("🏊‍♀️ La Lección de Vida", "No se aprende a nadar mirando la piscina; la maestría se construye ejecutando adentro.", c_teal)
    ]
    for i, (title, desc, color) in enumerate(cards9):
        card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.9), Inches(2.4), Inches(3.6), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = color
        card.line.width = Pt(2)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = c_white
        p.space_after = Pt(14)
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(19)
        p2.font.color.rgb = c_white

    # SLIDE 10
    build_pattern_slide(10, "DISCIPLINA vs. EMOCIÓN", "La Acción es Fría: Desmitificando la Motivación", [
        ("📉 La Falacia de la Motivación", "Esperar 'sentir ganas' te paraliza porque la emoción previa es voluble e inestable.", c_pink),
        ("🧊 La Acción es Neutra y Fría", "Accionar no requiere emoción previa; requiere decisión en automático sin pedir permiso a tu estado de ánimo.", c_teal),
        ("🏆 El Orden Real", "El orgullo, la satisfacción y la motivación vienen DESPUÉS de accionar, nunca antes.", c_white)
    ])

    # SLIDE 11: EJERCICIO
    s11 = prs.slides.add_slide(blank_layout)
    apply_bg(s11)
    add_header(s11, "PNL & PRÁCTICA EN VIVO", "Ejercicio Práctico: De la Teoría a la Acción")
    box11 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.5), Inches(10.333), Inches(3.8))
    box11.fill.solid()
    box11.fill.fore_color.rgb = c_card
    box11.line.color.rgb = c_teal
    box11.line.width = Pt(2)
    tf11 = box11.text_frame
    tf11.word_wrap = True
    tf11.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf11.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "🧘‍♀️ Ejercicio Práctico del Club"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = c_white
    p.space_after = Pt(16)
    p2 = tf11.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Guiado en vivo por Marbelis Montero"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = c_teal

    # SLIDE 12
    s12 = prs.slides.add_slide(blank_layout)
    apply_bg(s12)
    add_header(s12, "TRANSFORMACIÓN DE IDENTIDAD", "El Salto de Identidad y Compromiso")
    
    b12_1 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.2))
    b12_1.fill.solid()
    b12_1.fill.fore_color.rgb = c_card
    b12_1.line.color.rgb = c_teal
    tf = b12_1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👤 Cambio de Postura"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = c_teal
    p.space_after = Pt(6)
    p2 = tf.add_paragraph()
    p2.text = "Dejar de identificarte como 'la persona que analiza/sana' y encarnar a 'la persona que ejecuta con lo que tiene hoy'."
    p2.font.size = Pt(18)
    p2.font.color.rgb = c_white

    b12_2 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(1.8), Inches(5.8), Inches(2.2))
    b12_2.fill.solid()
    b12_2.fill.fore_color.rgb = c_card
    b12_2.line.color.rgb = c_pink
    tf = b12_2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🤝 Compromiso de 24 Horas"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = c_pink
    p.space_after = Pt(6)
    p2 = tf.add_paragraph()
    p2.text = "Ejecutar tu acción mínima hoy antes de irte a dormir, sin negociar con las excusas, y compartirla en el Club."
    p2.font.size = Pt(18)
    p2.font.color.rgb = c_white

    qbox12 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(11.7), Inches(2.4))
    qbox12.fill.solid()
    qbox12.fill.fore_color.rgb = c_card
    qbox12.line.color.rgb = c_teal
    qbox12.line.width = Pt(2)
    tf = qbox12.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "“No esperes a estar lista para lanzarte al agua; lánzate al agua para aprender a estar lista. Al final del día, tu potencial no se mide por lo impecable que es tu idea en la mente, sino por la valentía con la que encarnas tu nueva identidad en la realidad.”"
    p.font.size = Pt(22)
    p.font.italic = True
    p.font.color.rgb = c_white

    # SLIDE 13: EXACT LAYOUT FROM PREVIOUS SESSION
    s13 = prs.slides.add_slide(blank_layout)
    apply_bg(s13)
    add_header(s13, "COMUNIDAD EN VIVO", "Espacio de Preguntas & Reprogramación")
    
    box13 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(3.0), Inches(8.9), Inches(1.8))
    box13.fill.solid()
    box13.fill.fore_color.rgb = RGBColor(40, 15, 60)
    box13.line.color.rgb = c_teal
    box13.line.width = Pt(3)
    tf13 = box13.text_frame
    tf13.word_wrap = True
    tf13.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf13.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "⏸️ Pausamos la grabación para conversar"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = c_teal
    
    tb13_sub = s13.shapes.add_textbox(Inches(1.2), Inches(5.2), Inches(10.9), Inches(1.0))
    tf13_sub = tb13_sub.text_frame
    p = tf13_sub.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Un espacio seguro y privado para nuestra comunidad"
    p.font.size = Pt(22)
    p.font.color.rgb = c_white

    out_path = r"c:\Users\soyma\OneDrive\Documents\ANTIGRAVITY\MARCA\Realmente_Club\Diapositivas_Clase_Sabado_EP57.pptx"
    prs.save(out_path)
    print(f"PPTX updated successfully to {out_path}")

if __name__ == "__main__":
    build_presentation()
