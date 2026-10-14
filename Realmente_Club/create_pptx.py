import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # BRAND COLOR PALETTE: Deep Dark Grape Purple ("Morado Uva Oscuro")
    c_bg = RGBColor(25, 6, 51)           # #190633 - Deep Grape Dark Purple (Fondo Uva Oscuro)
    c_card = RGBColor(46, 12, 89)        # #2E0C59 - Plum Purple Card (Tarjetas Morado Ciruela)
    c_teal = RGBColor(42, 210, 197)      # #2AD2C5 - Teal Accent (Verde Agua Resaltante)
    c_pink = RGBColor(255, 90, 141)      # #FF5A8D - Magenta Pink Accent (Rosa Resaltante)
    c_lavender = RGBColor(212, 178, 230) # #D4B2E6 - Soft Lavender (Lila Suave)
    c_white = RGBColor(255, 255, 255)
    c_muted = RGBColor(220, 205, 235)    # #DCCDEB

    def apply_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_bg
        bg.line.fill.background()

    def add_header(slide, tag_text, title_text, subtitle_text=None):
        tb_tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_tag = tb_tag.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.size = Pt(13)
        p_tag.font.bold = True
        p_tag.font.color.rgb = c_teal

        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(0.8))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(32)
        p_title.font.bold = True
        p_title.font.color.rgb = c_white

        if subtitle_text:
            tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(0.5))
            tf_sub = tb_sub.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(18)
            p_sub.font.color.rgb = c_lavender

    # SLIDE 1: PORTADA PRINCIPAL (MORADO UVA OSCURO, CENTRADO)
    slide1 = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide1)
    
    tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    p_tag = tf1.paragraphs[0]
    p_tag.alignment = PP_ALIGN.CENTER
    p_tag.text = "REALMENTE CLUB · SESIÓN EN VIVO"
    p_tag.font.size = Pt(18)
    p_tag.font.bold = True
    p_tag.font.color.rgb = c_teal
    p_tag.space_after = Pt(20)

    p_title = tf1.add_paragraph()
    p_title.alignment = PP_ALIGN.CENTER
    p_title.text = "Reprogramación del Diálogo Interno"
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = c_white
    p_title.space_after = Pt(16)

    p_sub = tf1.add_paragraph()
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.text = "Del Jefe Tirano a la Verdad sin Juicio y los Pensamientos Constructivos"
    p_sub.font.size = Pt(22)
    p_sub.font.bold = True
    p_sub.font.color.rgb = c_lavender

    tb1_foot = slide1.shapes.add_textbox(Inches(1.0), Inches(6.5), Inches(11.333), Inches(0.6))
    tf1_foot = tb1_foot.text_frame
    p_foot = tf1_foot.paragraphs[0]
    p_foot.alignment = PP_ALIGN.CENTER
    p_foot.text = "Facilitadora: Marbelis Montero  |  Sábado 12 de Septiembre, 2026"
    p_foot.font.size = Pt(14)
    p_foot.font.color.rgb = c_muted

    # SLIDE 2: PLATON
    slide2 = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide2)
    add_header(slide2, "REFLEXIÓN DE APERTURA", "Tu Conversación Más Importante")
    
    qbox = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(2.2), Inches(10.9), Inches(2.6))
    qbox.fill.solid()
    qbox.fill.fore_color.rgb = c_card
    qbox.line.color.rgb = c_lavender
    qbox.line.width = Pt(2)
    tf2 = qbox.text_frame
    tf2.word_wrap = True
    tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf2.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "“El pensamiento es el diálogo que el alma sostiene consigo misma.”"
    p.font.size = Pt(28)
    p.font.italic = True
    p.font.color.rgb = c_white
    p.space_after = Pt(14)
    p_auth = tf2.add_paragraph()
    p_auth.alignment = PP_ALIGN.CENTER
    p_auth.text = "— PLATÓN"
    p_auth.font.size = Pt(22)
    p_auth.font.bold = True
    p_auth.font.color.rgb = c_teal
    
    qbox2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(5.2), Inches(10.9), Inches(1.5))
    qbox2.fill.solid()
    qbox2.fill.fore_color.rgb = RGBColor(60, 15, 65)
    qbox2.line.color.rgb = c_pink
    qbox2.line.width = Pt(2)
    tf2_2 = qbox2.text_frame
    tf2_2.word_wrap = True
    tf2_2.vertical_anchor = MSO_ANCHOR.MIDDLE
    qp1 = tf2_2.paragraphs[0]
    qp1.alignment = PP_ALIGN.CENTER
    qp1.text = "❓ Pregunta Disparadora para el Chat"
    qp1.font.size = Pt(20)
    qp1.font.bold = True
    qp1.font.color.rgb = c_pink
    qp1.space_after = Pt(6)
    qp2 = tf2_2.add_paragraph()
    qp2.alignment = PP_ALIGN.CENTER
    qp2.text = "¿Qué clase de diálogo has estado sosteniendo contigo esta semana?"
    qp2.font.size = Pt(22)
    qp2.font.bold = True
    qp2.font.color.rgb = c_white

    # SLIDE 3: UNIDAD DE MEDIDA
    slide3 = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide3)
    add_header(slide3, "CONCEPTO CLAVE DE LA CLASE", "Tu Diálogo Interno Nunca Miente", "La unidad de medida de la relación contigo misma")
    
    pillars = [
        ("📏 La Unidad de Medida", "La unidad de medida de la relación contigo misma es tu diálogo interno.", c_teal),
        ("💡 Sin Juicio ni Exigencia", "No se trata de ponerle expectativas irreales a tu proceso ni exigirte perfección.", c_pink),
        ("🩺 Radiografía Exacta", "Tu voz interna es la radiografía fiel de cómo te sientes y estás contigo en el presente.", c_lavender)
    ]
    for i, (title, desc, color) in enumerate(pillars):
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.9), Inches(2.6), Inches(3.6), Inches(4.2))
        card.fill.solid()
        card.fill.fore_color.rgb = c_card
        card.line.color.rgb = color
        card.line.width = Pt(2)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = c_white
        p.space_after = Pt(14)
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(17)
        p2.font.color.rgb = c_muted

    # SLIDE 4: ANGEL VS DIABLITO
    slide4 = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide4)
    add_header(slide4, "CONEXIÓN PROFUNDA Y ORIGEN DE LAS VOCES", "El Ángel vs. El Diablito y Tu Esencia")
    
    rows4 = [
        ("😇👿 Metáfora del Ángel y el Diablito (Las Comiquitas)", "En las caricaturas siempre se representaba la lucha entre la voz de la paz (angelito) y la voz del juicio/exigencia (diablito).", c_teal),
        ("📜 El Origen de la Lucha Interna", "Esa pelea interna no es casualidad; se activa desde lo más profundo de tus vivencias pasadas, tu historia y tus experiencias.", c_pink),
        ("✨ Tu Verdadera Esencia", "Detrás de la lucha entre esas dos voces habita quien realmente eres. Reconocerla te permite reconectar con tu verdad.", c_lavender)
    ]
    for i, (rtitle, rdesc, rcolor) in enumerate(rows4):
        box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2 + i * 1.6), Inches(11.7), Inches(1.3))
        box.fill.solid()
        box.fill.fore_color.rgb = c_card
        box.line.color.rgb = rcolor
        box.line.width = Pt(2)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = rtitle
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = rcolor
        p.space_after = Pt(4)
        p2 = tf.add_paragraph()
        p2.text = rdesc
        p2.font.size = Pt(17)
        p2.font.color.rgb = c_white

    # SLIDE 5: DECIRTE LA VERDAD
    slide5 = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide5)
    add_header(slide5, "SINCERAMIENTO RADICAL", "La Fuerza de Decirte la Verdad (Sin Juicio)")
    
    rows5 = [
        ("🛡️ Decirte la verdad NO es darte con el látigo", "No se trata de criticarte por fallar o exigirte estar siempre perfecta. La verdad nace de la compasión.", c_pink),
        ("💬 Admitir tu estado emocional real sin esconderlo", "Decirte con honestidad: “Hoy me siento cansada”, “Hoy siento miedo” o “Hoy no me siento capaz”.", c_teal),
        ("🕊️ La mente suelta la resistencia y entra en calma", "Al validar lo que sientes sin culpa ni vergüenza, la pelea interna se termina y el cuerpo descansa.", c_lavender)
    ]
    for i, (rtitle, rdesc, rcolor) in enumerate(rows5):
        box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2 + i * 1.6), Inches(11.7), Inches(1.3))
        box.fill.solid()
        box.fill.fore_color.rgb = c_card
        box.line.color.rgb = rcolor
        box.line.width = Pt(2)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = rtitle
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = rcolor
        p.space_after = Pt(4)
        p2 = tf.add_paragraph()
        p2.text = rdesc
        p2.font.size = Pt(17)
        p2.font.color.rgb = c_white

    # SLIDE 6: EJERCICIO PNL (SOLO TÍTULO CENTRADO GRANDE)
    slide6 = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide6)
    
    tb6 = slide6.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.2))
    tf6 = tb6.text_frame
    tf6.word_wrap = True
    tf6.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    p6 = tf6.paragraphs[0]
    p6.alignment = PP_ALIGN.CENTER
    p6.text = "🧘‍♀️ Ejercicio: Desarmando Tu Diálogo Interno"
    p6.font.size = Pt(42)
    p6.font.bold = True
    p6.font.color.rgb = c_white

    # SLIDE 7: MATRIZ I
    slide7 = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide7)
    add_header(slide7, "MATRIZ DE PENSAMIENTOS PUENTE · PARTE I", "De la Lucha al Pensamiento Constructivo")
    
    m1 = [
        ('"Tengo que hacer todo esto hoy o fallé"', '"Elijo hacer lo prioritario hoy y me permito avanzar a mi ritmo sin castigarme."'),
        ('"Debería estar más avanzada a mi edad"', '"Acepto dónde estoy hoy; mi proceso no se parece al de nadie más y está bien."'),
        ('"No soy capaz de resolver este problema"', '"Aún no tengo la respuesta, pero estoy en proceso de aprender a resolverlo."'),
        ('"Cometí un error horrible, qué tonta fui"', '"Cometí una falla y es un dato de aprendizaje; no define mi capacidad ni mi valor."')
    ]
    for i, (bad, good) in enumerate(m1):
        top = Inches(2.1 + i * 1.25)
        b1 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(5.6), Inches(1.15))
        b1.fill.solid()
        b1.fill.fore_color.rgb = c_card
        b1.line.color.rgb = c_pink
        tf1 = b1.text_frame
        tf1.word_wrap = True
        tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf1.paragraphs[0]
        p.text = bad
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 179, 198)

        b2 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), top, Inches(5.9), Inches(1.15))
        b2.fill.solid()
        b2.fill.fore_color.rgb = c_card
        b2.line.color.rgb = c_teal
        tf2 = b2.text_frame
        tf2.word_wrap = True
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf2.paragraphs[0]
        p.text = good
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(224, 251, 248)

    # SLIDE 8: MATRIZ II
    slide8 = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide8)
    add_header(slide8, "MATRIZ DE PENSAMIENTOS PUENTE · PARTE II", "De la Lucha al Pensamiento Constructivo")
    
    m2 = [
        ('"No me alcanza el tiempo para nada"', '"Decido organizar mi energía en lo que importa hoy y suelto la prisa."'),
        ('"Me da rabia/envidia que a ella le salga tan fácil"', '"Sinceramente siento rabia, y acepto que esto me muestra algo que yo también deseo."'),
        ('"No me siento lista para mostrarme"', '"No necesito sentirme 100% lista; me abro a dar el primer paso con lo que tengo hoy."')
    ]
    for i, (bad, good) in enumerate(m2):
        top = Inches(2.2 + i * 1.5)
        b1 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(5.6), Inches(1.35))
        b1.fill.solid()
        b1.fill.fore_color.rgb = c_card
        b1.line.color.rgb = c_pink
        tf1 = b1.text_frame
        tf1.word_wrap = True
        tf1.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf1.paragraphs[0]
        p.text = bad
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 179, 198)

        b2 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), top, Inches(5.9), Inches(1.35))
        b2.fill.solid()
        b2.fill.fore_color.rgb = c_card
        b2.line.color.rgb = c_teal
        tf2 = b2.text_frame
        tf2.word_wrap = True
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf2.paragraphs[0]
        p.text = good
        p.font.size = Pt(17)
        p.font.color.rgb = RGBColor(224, 251, 248)

    # SLIDE 9: LOUISE HAY
    slide9 = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide9)
    add_header(slide9, "REFLEXIÓN DE CIERRE", "Transformando la Relación Contigo")
    
    qbox9 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(2.2), Inches(10.9), Inches(2.8))
    qbox9.fill.solid()
    qbox9.fill.fore_color.rgb = c_card
    qbox9.line.color.rgb = c_lavender
    qbox9.line.width = Pt(2)
    tf9 = qbox9.text_frame
    tf9.word_wrap = True
    tf9.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    p = tf9.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "“Llevas años criticándote y no ha funcionado.\nPrueba a abordar la vida desde la aprobación y verás qué pasa.”"
    p.font.size = Pt(26)
    p.font.italic = True
    p.font.color.rgb = c_white
    p.space_after = Pt(14)
    
    p_auth = tf9.add_paragraph()
    p_auth.alignment = PP_ALIGN.CENTER
    p_auth.text = "— LOUISE HAY"
    p_auth.font.size = Pt(22)
    p_auth.font.bold = True
    p_auth.font.color.rgb = c_lavender

    box9_sub = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(5.4), Inches(10.9), Inches(1.3))
    box9_sub.fill.solid()
    box9_sub.fill.fore_color.rgb = RGBColor(40, 15, 60)
    box9_sub.line.color.rgb = c_teal
    box9_sub.line.width = Pt(2)
    tf9_sub = box9_sub.text_frame
    tf9_sub.word_wrap = True
    tf9_sub.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf9_sub.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "🌿 Sustituir la crítica dura por la aprobación amigable es el primer paso a la paz."
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RGBColor(224, 251, 248)

    # SLIDE 10: ESPACIO Q&A
    slide10 = prs.slides.add_slide(blank_slide_layout)
    apply_background(slide10)
    add_header(slide10, "COMUNIDAD EN VIVO", "Espacio de Preguntas & Reprogramación")
    
    box10 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(3.0), Inches(8.9), Inches(1.8))
    box10.fill.solid()
    box10.fill.fore_color.rgb = RGBColor(40, 15, 60)
    box10.line.color.rgb = c_teal
    box10.line.width = Pt(3)
    tf10 = box10.text_frame
    tf10.word_wrap = True
    tf10.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf10.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "⏸️ Pausamos la grabación para conversar"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = c_teal
    
    tb10_sub = slide10.shapes.add_textbox(Inches(1.2), Inches(5.2), Inches(10.9), Inches(1.0))
    tf10_sub = tb10_sub.text_frame
    p = tf10_sub.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Un espacio seguro y privado para nuestra comunidad"
    p.font.size = Pt(24)
    p.font.color.rgb = c_lavender

    out_path = r"c:\Users\soyma\OneDrive\Documents\ANTIGRAVITY\MARCA\Realmente_Club\Diapositivas_Clase_Sabado_Semana2.pptx"
    prs.save(out_path)
    print(f"Deep Grape Purple presentation saved successfully to {out_path}")

if __name__ == "__main__":
    build_presentation()
